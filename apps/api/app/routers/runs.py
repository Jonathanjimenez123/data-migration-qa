"""Endpoints para consultar y exportar el historial de ejecuciones de QA."""

import io
import json
from datetime import date, datetime
from decimal import Decimal
from uuid import UUID

from fastapi import APIRouter, HTTPException, Query
from fastapi.encoders import jsonable_encoder
from fastapi.responses import JSONResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.database import get_run, list_runs


router = APIRouter(prefix="/api/v1/runs", tags=["Execution History"])


def _serialize(value):
    """Convierte valores PostgreSQL a tipos compatibles con JSON."""
    return jsonable_encoder(
        value,
        custom_encoder={
            Decimal: float,
            datetime: lambda item: item.isoformat(),
            date: lambda item: item.isoformat(),
            UUID: str,
        },
    )


def _load_run(run_id: int):
    """Obtiene una ejecución y gestiona los errores de consulta."""
    if run_id < 1:
        raise HTTPException(
            status_code=422,
            detail="El identificador debe ser un entero positivo.",
        )

    try:
        run = get_run(run_id)
    except Exception as exc:
        raise HTTPException(
            status_code=503,
            detail="No fue posible consultar la ejecución.",
        ) from exc

    if run is None:
        raise HTTPException(
            status_code=404,
            detail="No se encontró esa ejecución.",
        )

    return run


def _excel_value(value):
    """Convierte estructuras complejas en texto compatible con Excel."""
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, Decimal):
        return float(value)
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _add_table_sheet(workbook, title, records):
    """Añade una hoja tabular sin asumir nombres fijos de columnas."""
    sheet = workbook.create_sheet(title)

    if not isinstance(records, list):
        records = []

    records = [
        record if isinstance(record, dict) else {"value": record}
        for record in records
    ]

    columns = []
    for record in records:
        for key in record:
            if key not in columns:
                columns.append(key)

    if not columns:
        sheet.append(["Resultado"])
        sheet.append(["No hay registros para mostrar."])
        return

    sheet.append(columns)

    for record in records:
        sheet.append([_excel_value(record.get(column)) for column in columns])

    _style_sheet(sheet)


def _style_sheet(sheet):
    """Aplica formato legible y ajusta el ancho de las columnas."""
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="17365D",
        )

    for column_cells in sheet.columns:
        max_length = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(value), 60))

        column_letter = get_column_letter(column_cells[0].column)
        sheet.column_dimensions[column_letter].width = max(
            12, min(max_length + 2, 62)
        )


@router.get("")
def read_runs(
    limit: int = Query(default=20, ge=1, le=100),
    offset: int = Query(default=0, ge=0),
):
    """Lista las ejecuciones más recientes."""
    try:
        runs = list_runs(limit=limit, offset=offset)
        return {
            "items": _serialize(runs),
            "limit": limit,
            "offset": offset,
            "count": len(runs),
        }
    except Exception as exc:
        raise HTTPException(
            status_code=503,
            detail="No fue posible consultar el historial de ejecuciones.",
        ) from exc


@router.get("/{run_id}/report.json")
def export_run_json(run_id: int):
    """Descarga el informe completo de una ejecución en JSON."""
    run = _serialize(_load_run(run_id))

    return JSONResponse(
        content=run,
        headers={
            "Content-Disposition": (
                f'attachment; filename="reporte_qa_ejecucion_{run_id}.json"'
            )
        },
    )


@router.get("/{run_id}/report.xlsx")
def export_run_excel(run_id: int):
    """Descarga el informe de una ejecución en Excel."""
    run = _serialize(_load_run(run_id))

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Resumen"
    summary_sheet.append(["Campo", "Valor"])

    for key, value in run.items():
        if key not in ("checks", "differences"):
            summary_sheet.append([key, _excel_value(value)])

    _style_sheet(summary_sheet)

    _add_table_sheet(
        workbook,
        "Comprobaciones",
        run.get("checks", []),
    )
    _add_table_sheet(
        workbook,
        "Diferencias",
        run.get("differences", []),
    )

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="reporte_qa_ejecucion_{run_id}.xlsx"'
            )
        },
    )


@router.get("/{run_id}")
def read_run(run_id: int):
    """Devuelve una ejecución con comprobaciones y diferencias."""
    return _serialize(_load_run(run_id))